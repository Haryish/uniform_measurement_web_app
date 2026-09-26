
from flask import Flask, render_template, request, send_file, flash, redirect, url_for
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path
from io import BytesIO
import random
import re
import tempfile
import os

app = Flask(__name__)
app.secret_key = "uniform-measurement-local-app"

UPLOAD_DIR = Path(tempfile.gettempdir()) / "uniform_measurement_uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

def clean(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())

def parse_group_from_sheet(title):
    t = clean(title)
    gender = "Boys" if "boy" in t or "male" in t else "Girls" if "girl" in t or "female" in t else None
    m = re.search(r"class\s*([0-9]+)(?:\s*-\s*([0-9]+))?", t)
    if not m:
        m = re.search(r"([0-9]+)\s*-\s*([0-9]+)", t)
    if m:
        c1=int(m.group(1)); c2=int(m.group(2) or c1)
        classes=list(range(c1,c2+1))
    else:
        classes=[]
    return gender, classes

def find_pattern_header(ws):
    for r in range(1, min(ws.max_row, 15)+1):
        vals=[clean(ws.cell(r,c).value) for c in range(1, ws.max_column+1)]
        if any(v=="pattern" or v.startswith("pattern") or "வடிவம்" in v for v in vals):
            return r
    return None

def read_patterns(fileobj):
    wb=load_workbook(fileobj, data_only=True)
    patterns=[]
    for ws in wb.worksheets:
        header_row=find_pattern_header(ws)
        if not header_row:
            continue
        headers=[clean(ws.cell(header_row,c).value) for c in range(1,ws.max_column+1)]
        pcol=next((i+1 for i,h in enumerate(headers) if h=="pattern" or h.startswith("pattern") or "வடிவம்" in h),1)
        gender, classes=parse_group_from_sheet(ws.title)
        if not gender:
            # infer from pattern IDs
            sample=str(ws.cell(header_row+1,pcol).value or "")
            gender="Girls" if sample.upper().startswith("G") else "Boys"
        if not classes:
            # infer class from pattern ID, e.g. B3-1 or G5-7-1
            for rr in range(header_row+1, min(ws.max_row,header_row+4)+1):
                s=str(ws.cell(rr,pcol).value or "")
                nums=re.findall(r"\d+",s)
                if nums:
                    classes=[int(nums[0])]
                    break
        for r in range(header_row+1, ws.max_row+1):
            pid=ws.cell(r,pcol).value
            if pid in (None,""): continue
            vals=[]
            for c in range(pcol+1,ws.max_column+1):
                v=ws.cell(r,c).value
                if isinstance(v,(int,float)) and not isinstance(v,bool):
                    v=int(v)
                    if v <= 0 or v >= 100 or v % 2:
                        v = v + 1 if v > 0 and v < 99 and v % 2 else None
                    if v is not None: vals.append(v)
            if vals:
                patterns.append({"id":str(pid),"gender":gender,"classes":classes,"values":vals})
    return patterns

def even_value(v):
    if v is None or v == "":
        return None
    try:
        x=int(float(v))
    except Exception:
        return None
    if x <= 0 or x >= 100:
        return None
    return x if x%2==0 else x+1

def generate_patterns(base_patterns, variants=3, seed=42):
    rng=random.Random(seed)
    out=[]
    seen=set()
    for p in base_patterns:
        out.append(dict(p, values=[even_value(v) for v in p["values"]]))
        seen.add((p["gender"],tuple(p["classes"]),tuple(p["values"])))
    for p in base_patterns:
        for k in range(variants):
            vals=[]
            for v in p["values"]:
                delta=rng.choice([-4,-2,0,2,4])
                nv=max(10,min(98, v+delta))
                if nv%2: nv+=1
                vals.append(nv)
            key=(p["gender"],tuple(p["classes"]),tuple(vals))
            if key in seen: continue
            seen.add(key)
            out.append({
                "id":f'{p["id"]}-V{k+1}',
                "gender":p["gender"],
                "classes":p["classes"],
                "values":vals
            })
    return out

def patterns_to_workbook(patterns):
    wb=Workbook()
    wb.remove(wb.active)
    groups={}
    for p in patterns:
        for cls in p["classes"]:
            key=(p["gender"],cls)
            groups.setdefault(key,[]).append(p)
    for (gender,cls), plist in sorted(groups.items(), key=lambda x:(x[0][0],x[0][1])):
        ws=wb.create_sheet(f"{gender} - Class {cls}")
        maxn=max(len(p["values"]) for p in plist)
        headers=["Pattern"]+[f"Measurement {i}" for i in range(1,maxn+1)]
        ws.append(headers)
        for p in plist:
            ws.append([p["id"]]+p["values"])
        for cell in ws[1]:
            cell.font=Font(bold=True)
            cell.fill=PatternFill("solid", fgColor="D9EAF7")
            cell.alignment=Alignment(horizontal="center",vertical="center")
        for col in range(1,ws.max_column+1):
            ws.column_dimensions[chr(64+col) if col<=26 else "A"].width=16
    return wb

def find_template_header(ws):
    for r in range(1,min(ws.max_row,15)+1):
        vals=[clean(ws.cell(r,c).value) for c in range(1,ws.max_column+1)]
        if any(v=="s.no" or v=="s. no." for v in vals) and any("student" in v for v in vals):
            return r
    return None

def template_info(ws):
    hr=find_template_header(ws)
    if not hr:
        raise ValueError(f"Could not find the student table header in sheet '{ws.title}'.")
    headers=[clean(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
    section_col=next((i+1 for i,h in enumerate(headers) if h=="section"),6)
    measurement_cols=list(range(section_col+1,ws.max_column+1))
    class_col=next((i+1 for i,h in enumerate(headers) if h=="class"),5)
    gender_col=next((i+1 for i,h in enumerate(headers) if h=="gender"),3)
    name_col=next((i+1 for i,h in enumerate(headers) if "student" in h),2)
    return hr, measurement_cols, class_col, gender_col, name_col

def choose_patterns(all_patterns, gender, cls, allow_fallback=True):
    exact=[p for p in all_patterns if p["gender"]==gender and cls in p["classes"]]
    if exact: return exact, False
    if allow_fallback:
        same=[p for p in all_patterns if p["gender"]==gender]
        if same:
            # nearest class group
            def dist(p):
                return min(abs(cls-c) for c in p["classes"])
            d=min(dist(p) for p in same)
            return [p for p in same if dist(p)==d], True
    return [], False

def fill_template(template_file, patterns, mode="round_robin", seed=42, allow_fallback=True, overwrite=False):
    wb=load_workbook(template_file)
    rng=random.Random(seed)
    report=[]
    for ws in wb.worksheets:
        try:
            hr, mcols, class_col, gender_col, name_col=template_info(ws)
        except ValueError as e:
            report.append(("SKIPPED",ws.title,str(e)))
            continue
        counters={}
        for r in range(hr+1, ws.max_row+1):
            name=ws.cell(r,name_col).value
            if not name: continue
            raw_gender=str(ws.cell(r,gender_col).value or "")
            gender="Girls" if "female" in raw_gender.lower() or "girl" in raw_gender.lower() else "Boys"
            try: cls=int(str(ws.cell(r,class_col).value).strip())
            except: 
                report.append(("WARNING",ws.title,f"Row {r}: invalid class"))
                continue
            plist,fallback=choose_patterns(patterns,gender,cls,allow_fallback)
            if not plist:
                report.append(("WARNING",ws.title,f"Row {r} ({name}): no pattern for {gender} class {cls}"))
                continue
            key=(gender,cls)
            idx=counters.get(key,0)
            if mode=="random":
                p=rng.choice(plist)
            else:
                p=plist[idx % len(plist)]
            counters[key]=idx+1
            vals=p["values"]
            if len(vals) < len(mcols):
                report.append(("WARNING",ws.title,f"Row {r} ({name}): pattern {p['id']} has {len(vals)} values but template expects {len(mcols)}"))
            for i,c in enumerate(mcols):
                if i>=len(vals): break
                if overwrite or ws.cell(r,c).value in (None,""):
                    ws.cell(r,c).value=even_value(vals[i])
            report.append(("FILLED",ws.title,f"Row {r}: {name} <- {p['id']}"+(" (fallback)" if fallback else "")))
    return wb,report

@app.route("/", methods=["GET","POST"])
def index():
    if request.method=="POST":
        pattern_file=request.files.get("pattern_file")
        template_file=request.files.get("template_file")
        mode=request.form.get("mode","round_robin")
        seed=int(request.form.get("seed","42") or 42)
        allow_fallback=request.form.get("fallback")=="on"
        overwrite=request.form.get("overwrite")=="on"
        if not pattern_file or not template_file:
            flash("Please select both a pattern workbook and a template workbook.")
            return redirect(url_for("index"))
        try:
            patterns=read_patterns(pattern_file)
            if not patterns:
                raise ValueError("No pattern sheets were detected. A sheet must contain a 'Pattern' heading and numeric measurement columns.")
            wb,report=fill_template(template_file,patterns,mode,seed,allow_fallback,overwrite)
            out=BytesIO()
            wb.save(out); out.seek(0)
            session_path=UPLOAD_DIR/"completed_uniform_measurements.xlsx"
            with open(session_path,"wb") as f: f.write(out.getbuffer())
            filled=sum(1 for x in report if x[0]=="FILLED")
            warnings=sum(1 for x in report if x[0]=="WARNING")
            return render_template("result.html", filled=filled,warnings=warnings,report=report[-100:])
        except Exception as e:
            flash(str(e))
            return redirect(url_for("index"))
    return render_template("index.html")

@app.route("/download")
def download():
    p=UPLOAD_DIR/"completed_uniform_measurements.xlsx"
    if not p.exists():
        flash("No generated file is available yet.")
        return redirect(url_for("index"))
    return send_file(p,as_attachment=True,download_name="completed_uniform_measurements.xlsx")

@app.route("/generate-patterns", methods=["POST"])
def generate():
    pattern_file=request.files.get("pattern_file")
    variants=int(request.form.get("variants","3") or 3)
    seed=int(request.form.get("seed","42") or 42)
    if not pattern_file:
        flash("Upload an existing pattern workbook first.")
        return redirect(url_for("index"))
    try:
        base=read_patterns(pattern_file)
        if not base: raise ValueError("No patterns found.")
        generated=generate_patterns(base,variants,seed)
        wb=patterns_to_workbook(generated)
        out=BytesIO(); wb.save(out); out.seek(0)
        return send_file(out,as_attachment=True,download_name="generated_patterns.xlsx")
    except Exception as e:
        flash(str(e)); return redirect(url_for("index"))

if __name__=="__main__":
    app.run(host="127.0.0.1",port=5000,debug=True)
